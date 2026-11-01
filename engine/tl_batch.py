"""Timeline Collision from a spreadsheet: one row = one video (country comparisons or any number line).

Sheet layout (first row = headers, English or Arabic, any order; see write_template()):
    Title | Country 1 | Value 1 | … | Country 5 | Value 5 | Hook | Reveal | Intro | Map style | Guess |
    Ending question | (optional) Unit | Mode | Name 1-5 | Color 1-5 | Icon 1-5 | Maps | Doubt | Doubt subtitle |
    Compare title | Guess question | Ending intro | CTA | Theme | Music

- Country: Arabic or English name, or the 3-letter code (EGY). Leave it empty and fill Name for a non-country item.
- Value: the number (area, height…); "2,149,690" and Arabic digits are fine.
- Empty cells use the editor's settings; an empty Hook / Reveal is written automatically for comparisons.
- New lines inside a cell (Alt+Enter) or " | " split text into lines; [gold] at a line start colours it.
"""
import copy
import csv
import os
import re

from . import countries as CT
from .batch import _norm
from .core import COLORS, THEME_NAMES
from .strings import ratio_str

MAX_EVENTS = 5
AREA_UNIT = {"ar": "كم²", "en": "km²"}
AUTO_TEXT = {"ar": {"area": "تابع المساحة…", "numbers": "تابع الأرقام…", "true_size": "بالحجم الحقيقي…"},
             "en": {"area": "Watch the area…", "numbers": "Watch the numbers…", "true_size": "True size…"}}
PALETTE = ["red", "green", "blue", "gold", "pink", "orange", "purple", "cyan", "sand"]

_ALIASES = {
    "title": ["title", "videotitle", "video", "العنوان", "اسمالفيديو", "عنوانالفيديو"],
    "mode": ["mode", "type", "النوع", "الوضع", "النمط"],
    "unit": ["unit", "units", "الوحده", "وحده"],
    "hook": ["hook", "claim", "الافتتاحيه", "الجملهالافتتاحيه", "الخطاف", "المقدمه"],
    "doubt": ["doubt", "doubtword", "التشكيك", "كلمهالتشكيك"],
    "doubt_sub": ["doubtsubtitle", "doubtsub", "سطرالتشكيك", "تحتالتشكيك"],
    "compare_title": ["comparetitle", "comparisontitle", "compare", "عنوانالمقارنه"],
    "reveal": ["reveal", "result", "الكشف", "النتيجه"],
    "intro": ["intro", "introstyle", "start", "نوعالبدايه", "البدايه"],
    "map_style": ["mapstyle", "mapscene", "style", "شكلالخريطه", "نمطالخريطه", "شكلالخرائط"],
    "maps": ["maps", "map", "truesize", "الخرائط", "خرائط", "الخريطه"],
    "guess": ["guess", "التخمين", "خمن"],
    "guess_q": ["guessquestion", "سؤالالتخمين"],
    "end_intro": ["endingintro", "endintro", "مقدمهالنهايه"],
    "end_q": ["endingquestion", "nextquestion", "ending", "سؤالالنهايه", "سؤالالحلقهالقادمه"],
    "cta": ["cta", "calltoaction", "الدعوه", "دعوه"],
    "theme": ["theme", "الثيم", "السمه"],
    "music": ["music", "الموسيقي", "موسيقي"],
}
_PER_EVENT = {
    "country": ["country", "دوله", "الدوله", "البلد", "بلد"],
    "value": ["value", "area", "year", "number", "القيمه", "المساحه", "السنه", "الرقم"],
    "name": ["name", "item", "اسم", "الاسم"],
    "color": ["color", "colour", "لون", "اللون"],
    "icon": ["icon", "ايقونه", "الايقونه"],
}
_LOOKUP = {_norm(a): k for k, names in _ALIASES.items() for a in names}
for _k, _names in _PER_EVENT.items():
    for _i in range(1, MAX_EVENTS + 1):
        for _a in _names:
            _LOOKUP[_norm(f"{_a}{_i}")] = f"{_k}{_i}"

_COUNTRY_ALIASES = {"usa": "USA", "us": "USA", "america": "USA", "امريكا": "USA", "الولاياتالمتحدهالامريكيه": "USA",
                    "uk": "GBR", "britain": "GBR", "greatbritain": "GBR", "unitedkingdom": "GBR", "بريطانيا": "GBR",
                    "المملكهالمتحده": "GBR", "uae": "ARE", "emirates": "ARE", "الامارات": "ARE", "ksa": "SAU",
                    "russia": "RUS", "southkorea": "KOR", "northkorea": "PRK", "كورياالجنوبيه": "KOR",
                    "drc": "COD", "congo": "COG", "czechia": "CZE", "تشيكيا": "CZE", "عمان": "OMN", "oman": "OMN",
                    "turkiye": "TUR", "سوريه": "SYR"}
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫", "01234567890123456789.")


# ------------------------------------------------------------------ reading
def _text(v):
    """Cell text keeping line breaks; ' | ' also starts a new line."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").replace(" | ", "\n")
    return "\n".join(" ".join(line.split()) for line in s.split("\n")).strip()


def read_rows(path):
    """Rows of the first sheet (.xlsx) or of a CSV file, as lists of strings (line breaks kept)."""
    ext = os.path.splitext(str(path))[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError as e:
            raise RuntimeError("Reading Excel needs the 'openpyxl' package: run run_windows.bat again "
                               "(it installs it) or `pip install openpyxl`.") from e
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [[_text(v) for v in row] for row in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    if ext in (".csv", ".txt", ".tsv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            text = f.read()
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        import io
        return [[_text(v) for v in row] for row in csv.reader(io.StringIO(text), dialect)]
    if ext == ".xls":
        raise ValueError("Old .xls files are not supported — in Excel use File › Save As › Excel Workbook (.xlsx).")
    raise ValueError(f"Unsupported file type '{ext}'. Use .xlsx or .csv.")


_country_index = None


def find_country(text):
    """Country code from an Arabic/English name or a code ('' when unknown)."""
    global _country_index
    if _country_index is None:
        idx = {}
        for code, c in CT._load().items():
            idx[code.lower()] = code
            for nm in (c["en"], c["ar"]):
                n = _norm(nm)
                idx[n] = code
                if n.startswith("ال"):
                    idx[n[2:]] = code  # "سعوديه" finds "السعوديه"
        idx.update(_COUNTRY_ALIASES)
        _country_index = idx
    n = _norm(re.sub(r"[\u064B-\u0652\u0640]", "", str(text or "")))  # ignore diacritics and tatweel
    for cand in (n, n[2:] if n.startswith("ال") else "", n.replace("غ", "ج")):  # غرينلاند = جرينلاند
        if cand and cand in _country_index:
            return _country_index[cand]
    return ""


def _number(text):
    s = str(text or "").translate(_DIGITS).replace(",", "").replace("،", "").replace(" ", "")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else None


def _bool(text, default):
    n = _norm(text)
    if n in ("yes", "y", "true", "1", "on", "نعم", "اي", "ايوه", "صح"):
        return True
    if n in ("no", "n", "false", "0", "off", "لا", "خطا"):
        return False
    return default


def _choice(text, options, default):
    """options: {value: [aliases]}"""
    n = _norm(text)
    if not n:
        return default
    for val, names in options.items():
        if n == _norm(val) or n in (_norm(a) for a in names):
            return val
    return None


_MODES = {"scale": ["numberline", "number", "value", "values", "area", "رقم", "ارقام", "مساحه", "قيم"],
          "calendar": ["calendar", "years", "year", "date", "سنوات", "تقويم", "تاريخ"],
          "ago": ["yearsago", "deeptime", "منذ", "قبل"]}
_INTROS = {"hook": ["text", "hooktext", "normal", "نص", "عادي"], "vs": ["versus", "battle", "vsbattle", "مواجهه", "ضد"]}
_STYLES = {"auto": ["automatic", "تلقائي"], "overlay": ["over", "فوق", "تراكب"],
           "tiles": ["tile", "tilefill", "fill", "تعبئه", "تبليط", "ملء"],
           "side": ["sidebyside", "side by side", "جنبالجنب", "جنبا", "بجانب"]}


def parse(rows):
    """Turn sheet rows into video specs. Returns (specs, problems); rows with problems are skipped."""
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return [], ["The sheet is empty."]
    cols = {}
    for i, h in enumerate(rows[0]):
        key = _LOOKUP.get(_norm(h))
        if key and key not in cols:
            cols[key] = i
    if not any(k.startswith("value") for k in cols):
        return [], ["No 'Value 1' / 'Value 2' columns found — keep the template's header row."]
    specs, problems = [], []
    for rn, row in enumerate(rows[1:], 2):
        get = lambda k: row[cols[k]].strip() if k in cols and cols[k] < len(row) and row[cols[k]] else ""  # noqa: E731
        evs, bad = [], []
        for i in range(1, MAX_EVENTS + 1):
            c, v, nm = get(f"country{i}"), get(f"value{i}"), get(f"name{i}")
            if not (c or v or nm):
                continue
            code = find_country(c) if c else ""
            if c and not code:
                bad.append(f"country {i} '{c}' not found")
                continue
            val = _number(v)
            if val is None:
                bad.append(f"item {i} ({c or nm}) has no number in Value {i}")
                continue
            if not code and not nm:
                bad.append(f"item {i} needs a Country or a Name")
                continue
            evs.append(dict(country=code, name=nm, typed=c, value=val, color=get(f"color{i}"), icon=get(f"icon{i}")))
        if not evs and not bad:
            continue  # empty row
        if len(evs) < 2:
            bad.append("needs at least 2 items")
        if len({e["value"] for e in evs}) < len(evs):
            bad.append("two items have the same value")
        codes = [e["country"] for e in evs if e["country"]]
        if len(set(codes)) < len(codes):
            bad.append("the same country is used twice")
        mode = _choice(get("mode"), _MODES, "scale")
        intro = _choice(get("intro"), _INTROS, "hook")
        style = _choice(get("map_style"), _STYLES, "auto")
        for label, val, raw in (("Mode", mode, get("mode")), ("Intro", intro, get("intro")),
                                ("Map style", style, get("map_style"))):
            if val is None:
                bad.append(f"{label} '{raw}' is not one of the choices")
        if bad:
            problems.append(f"Row {rn}: " + "; ".join(bad))
            continue
        specs.append(dict(row=rn, events=evs, mode=mode, intro=intro, map_style=style,
                          unit=get("unit"), maps=get("maps"), guess=get("guess"), title=get("title"),
                          text={k: get(k) for k in ("hook", "doubt", "doubt_sub", "compare_title", "reveal",
                                                    "guess_q", "end_intro", "end_q", "cta")},
                          theme=get("theme"), music=get("music")))
    return specs, problems


def load_sheet(path):
    return parse(read_rows(path))


# ------------------------------------------------------------------ building projects
def _flag_colors(code):
    """Palette colours ordered by how close they are to the flag's most vivid colour."""
    fp = CT.flag_path(code)
    if not fp:
        return PALETTE
    from PIL import Image
    im = Image.open(fp).convert("RGB").resize((48, 32))
    best, score = None, -1.0
    for r, g, b in im.getdata():
        mx, mn = max(r, g, b), min(r, g, b)
        sat = (mx - mn) / mx if mx else 0
        if sat > 0.35 and mx > 60 and sat * mx > score:
            best, score = (r, g, b), sat * mx
    if best is None:
        return PALETTE
    dist = lambda name: sum((a - b) ** 2 for a, b in zip(COLORS[name], best))  # noqa: E731
    return sorted(PALETTE, key=dist)


def display_name(e, lang):
    """Name column, else the country as typed in the sheet (when written in the video's language), else the
    country's name from the data (e.g. 'الإمارات' stays short instead of 'الإمارات العربية المتحدة')."""
    if e["name"]:
        return e["name"]
    typed = e.get("typed", "")
    arabic = bool(re.search(r"[\u0600-\u06FF]", typed))
    if typed and (arabic if lang == "ar" else (not arabic and _norm(typed) != e["country"].lower())):
        return typed
    return CT.name(e["country"], lang)


def _names(spec, lang):
    return [display_name(e, lang) for e in spec["events"]]


def auto_texts(spec, lang):
    """Hook and reveal written from the numbers (used when those cells are empty)."""
    evs = sorted(zip(spec["events"], _names(spec, lang)), key=lambda x: x[0]["value"])
    (small, a), (big, b) = evs[0], evs[-1]
    if small["value"] <= 0:
        return {}
    r = big["value"] / small["value"]
    times = ratio_str(r)[2:]
    pct = (r - 1) * 100
    pct_txt = f"{pct:.1f}".rstrip("0").rstrip(".") if pct < 10 else str(int(round(pct)))
    if lang == "ar":
        hook = f"{a} أم {b}؟\n[gold]أيهما أكبر؟"
        reveal = (f"[gold]الفرق بين {a} و{b}\n{pct_txt}% فقط!" if r < 1.5 else f"[gold]{b} أكبر من {a}\nبـ {times} مرة!")
    else:
        hook = f"{a} or {b}?\n[gold]Which is bigger?"
        reveal = (f"[gold]{a} and {b}:\nonly {pct_txt}% apart!" if r < 1.5 else f"[gold]{b} is {times}×\nthe size of {a}!")
    return {"hook": hook, "reveal": reveal}


def to_project(base, spec):
    """A normal Timeline Collision project for one row, built on the editor's settings."""
    p = copy.deepcopy(base)
    s, tl, end = p.setdefault("settings", {}), p.setdefault("timeline", {}), p.setdefault("ending", {})
    lang = s.get("language") or "ar"
    has_country = any(e["country"] for e in spec["events"])
    used, events = set(), []
    for e in spec["events"]:
        col = e["color"]
        if not col:
            order = _flag_colors(e["country"]) if e["country"] else PALETTE
            col = next((c for c in order if c not in used), order[0])
        used.add(col)
        icon = e["icon"].lower() if e["icon"] else ("flag" if e["country"] else "none")
        events.append(dict(enabled=True, country=e["country"], name=display_name(e, lang), year=e["value"], label="",
                           color=col, icon=icon, image="", outline=""))
    tl["events"] = events
    tl["mode"] = spec["mode"]
    tl["unit"] = spec["unit"] or (AREA_UNIT.get(lang, "km²") if has_country and spec["mode"] == "scale"
                                  else tl.get("unit", ""))
    n_countries = sum(1 for e in spec["events"] if e["country"])
    tl["maps"] = _bool(spec["maps"], n_countries >= 2 and spec["mode"] == "scale")
    tl["guess"] = _bool(spec["guess"], True)
    tl["intro"], tl["map_style"] = spec["intro"], spec["map_style"]
    tl["hide_first_in_hook"] = False
    auto = auto_texts(spec, lang) if spec["mode"] == "scale" else {}
    if spec["mode"] == "scale" and base.get("timeline", {}).get("mode") != "scale":
        # the editor holds a years project: its doubt subtitle / comparison title talk about time
        auto["doubt_sub"] = AUTO_TEXT[lang]["area" if has_country else "numbers"]
        if tl["maps"]:
            auto["compare_title"] = AUTO_TEXT[lang]["true_size"]
    txt = spec["text"]
    for key in ("hook", "doubt", "doubt_sub", "compare_title", "reveal", "guess_q"):
        if txt[key]:
            tl[key] = txt[key]
        elif key in auto:
            tl[key] = auto[key]
    if not txt["guess_q"]:
        tl["guess_q"] = ""  # automatic question
    for key, dst in (("end_intro", "intro"), ("end_q", "question"), ("cta", "cta")):
        if txt[key]:
            end[dst] = txt[key]
    if spec["theme"]:
        match = next((t for t in THEME_NAMES if _norm(t) == _norm(spec["theme"])), None)
        if match:
            s["theme"] = match
    if spec["music"]:
        s["music"] = spec["music"]
    return p


def file_name(spec, k, lang="ar"):
    base = spec["title"] or "_".join(_names(spec, lang))
    base = re.sub(r'[\\/:*?"<>|\n\r\t]+', " ", base).strip()[:60] or "video"
    return f"{k:02d}_{base.replace(' ', '_')}"


def summary(specs, problems):
    if not specs:
        txt = "⚠️ No usable rows."
    else:
        txt = f"**{len(specs)} video{'s' if len(specs) != 1 else ''}** ready (rows " \
              f"{', '.join(str(s['row']) for s in specs[:12])}{'…' if len(specs) > 12 else ''})."
    if problems:
        more = f"  \n…and {len(problems) - 30} more" if len(problems) > 30 else ""
        txt += "\n\n" + "  \n".join("⚠️ " + p for p in problems[:30]) + more
    return txt


# ------------------------------------------------------------------ template
HEADERS = ["Title", "Country 1", "Value 1", "Country 2", "Value 2", "Country 3", "Value 3", "Country 4", "Value 4",
           "Country 5", "Value 5", "Hook", "Reveal", "Intro", "Map style", "Guess", "Ending question", "Unit", "Mode",
           "Theme", "Music"]

EXAMPLES = [
    ("لبنان داخل السعودية", [("لبنان", 10452), ("قطر", 11586), ("السعودية", 2149690)],
     "لبنان يدخل في\n[green]السعودية\n[red]أكثر من 200 مرة!", "[red]لبنان يدخل في السعودية\n206 مرات!", "vs", "auto",
     "yes", "أيهما أكبر:\n[gold]غرينلاند أم السعودية؟"),
    ("غرينلاند أم السعودية", [("غرينلاند", 2166086), ("السعودية", 2149690)],
     "غرينلاند تبدو على الخرائط\n[blue]أكبر من أفريقيا…\n[green]لكنها بحجم السعودية تقريبًا!",
     "[gold]الفرق بينهما\nأقل من 1%!", "hook", "auto", "no", "مصر أقرب في المساحة\n[gold]إلى من؟"),
    ("مصر وموريتانيا والسعودية", [("مصر", 1001450), ("موريتانيا", 1030700), ("السعودية", 2149690)],
     "مصر أقرب في المساحة إلى\n[gold]موريتانيا\n[green]منها إلى السعودية!", "[green]السعودية أكبر من\nضعف مساحة مصر!",
     "hook", "auto", "yes", "كم بحرين تدخل\n[gold]في الكويت؟"),
    ("البحرين داخل الكويت", [("البحرين", 787), ("الكويت", 17818)], "", "", "vs", "auto", "yes",
     "الأردن أم البرتغال:\n[gold]أيهما أكبر؟"),
    ("الأردن أم البرتغال", [("الأردن", 89342), ("البرتغال", 92212)], "", "", "hook", "auto", "no",
     "غانا أم بريطانيا:\n[gold]أيهما أكبر؟"),
    ("غانا أم بريطانيا", [("غانا", 238535), ("بريطانيا", 242495)], "", "", "vs", "auto", "no",
     "أكبر 3 دول عربية:\n[gold]من الأولى؟"),
    ("أكبر الدول العربية", [("السودان", 1886068), ("السعودية", 2149690), ("الجزائر", 2381741)],
     "أكبر 3 دول عربية\n[gold]متقاربة جدًا في المساحة!", "[gold]الجزائر الأولى\nبفارق صغير!", "hook", "auto", "no",
     "كم كويت تدخل\n[gold]في العراق؟"),
    ("الكويت داخل العراق", [("الكويت", 17818), ("العراق", 438317)], "", "", "hook", "auto", "yes",
     "كم تونس تدخل\n[gold]في ليبيا؟"),
    ("تونس وليبيا", [("تونس", 163610), ("ليبيا", 1759540)], "", "", "vs", "auto", "yes",
     "الإمارات وعُمان واليمن:\n[gold]رتّبها حسب المساحة!"),
    ("الإمارات وعمان واليمن", [("الإمارات", 83600), ("عمان", 309500), ("اليمن", 527968)], "", "", "hook", "auto",
     "yes", "اكتب دولتين\n[gold]لنقارنهما!"),
]

NOTES = [
    "One row = one video. Keep the header row (English or Arabic headers both work, in any order).",
    "Country 1-5: Arabic or English name (مصر / Egypt) or the 3-letter code (EGY). 2 to 5 countries per row.",
    "Value 1-5: the official area in km² (or any number: heights, prices…). 2,149,690 or ٢١٤٩٦٩٠ are fine.",
    "Hook / Reveal: new lines with Alt+Enter (or ' | '). Start a line with [gold], [green], [red]… to colour it.",
    "   Leave Hook / Reveal empty and they are written automatically from the numbers.",
    "Intro: hook (text) or vs (flags battle).  Map style: auto, overlay, tiles, side.  Guess: yes / no.",
    "   auto = tile fill when one country is 20× bigger or more, side by side within 1.5×, overlay otherwise.",
    "Ending question: the question for the next video (tease the next row!).",
    "Optional: Unit (default كم²), Mode (number line / calendar years), Theme (e.g. Ocean), Music.",
    "Not a country? Add a 'Name 1' column (and leave Country 1 empty), e.g. برج خليفة with Value 828 and Unit متر.",
    "Everything else (handle, music, font, ending call-to-action…) comes from the app's 🎨 Style & sound and 🔚 Ending.",
]


def write_template(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Videos"
    ws.append(HEADERS)
    for title, items, hook, reveal, intro, style, guess, end_q in EXAMPLES:
        cells = []
        for i in range(MAX_EVENTS):
            cells += list(items[i]) if i < len(items) else ["", ""]
        ws.append([title, *cells, hook, reveal, intro, style, guess, end_q, "", "", "", ""])
    head = PatternFill("solid", fgColor="1F3A68")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(horizontal="center", vertical="center")
    widths = [22] + [14, 12] * MAX_EVENTS + [34, 30, 8, 11, 8, 30, 8, 10, 10, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i in range(MAX_EVENTS):
            row[2 + 2 * i].number_format = "#,##0.###"
    ws.freeze_panes = "B2"
    ws.sheet_view.rightToLeft = True
    notes = wb.create_sheet("How to fill")
    for line in NOTES:
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    wb.save(path)
    return path
